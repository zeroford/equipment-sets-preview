#!/usr/bin/env python3
"""สร้าง .xlsx จาก TSV ที่ sheet-seed.mjs export ไว้ — สำหรับ import เข้า Google Sheets ทีเดียวจบ

    ./render.sh && node tools/sheet-seed.mjs && python3 tools/sheet-xlsx.py
    → google-apps-script/seed/equipment-sets.xlsx

NOTE: ช่อง stat แบบ % เก็บเป็นเศษส่วน (0.0853) แล้วตั้ง number_format เป็น 0.00%
      เปิดในชีตจะเห็น 8.53% และพิมพ์ทับเป็น 8.53% ได้ตรงๆ
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SEED = Path(__file__).resolve().parent.parent / 'google-apps-script' / 'seed'
TABS = ['Sets', 'Items', 'Stats', 'BestStats']
OUT = SEED / 'equipment-sets.xlsx'

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(color='FFFFFF', bold=True)


def as_number(text):
    """ช่องที่เป็นตัวเลขล้วนให้เขียนเป็นตัวเลข ที่เหลือ (setKey, statId, ชื่อของ) เป็นข้อความ"""
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return text

wb = Workbook()
wb.remove(wb.active)

percent_cells = {tuple(c) for c in json.loads((SEED / 'percent-cells.json').read_text())}

for tab in TABS:
    rows = [line.split('\t') for line in (SEED / f'{tab}.tsv').read_text().rstrip('\n').split('\n')]
    ws = wb.create_sheet(tab)
    widths = {}

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value if r == 1 else as_number(value))
            if tab == 'Items' and (r, c) in percent_cells:
                cell.number_format = '0.00%'
            if r == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal='center')
            widths[c] = max(widths.get(c, 0), len(value))

    for c, width in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = min(max(width + 2, 8), 34)
    ws.freeze_panes = 'A2'

wb.save(OUT)
print(f'{OUT}  ({", ".join(f"{t}={wb[t].max_row - 1} rows" for t in TABS)})')
